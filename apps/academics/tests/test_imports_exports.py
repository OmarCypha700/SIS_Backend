"""CSV/Excel imports and exports, marks import, course and staff imports, and PDF documents."""
import io
import re

import pytest
from django.core import mail
from openpyxl import Workbook, load_workbook

from apps.academics.models import Course, ProgrammeCourse
from apps.accounts.models import Role, User
from apps.audit.models import AuditLog
from apps.registration import services as registration
from apps.results import services as results
from apps.results.models import StudentMark
from common.tabular import read_table
from conftest import client_for, make_user

pytestmark = pytest.mark.django_db


def xlsx_upload(rows, name="data.xlsx"):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = name
    return buffer


def csv_upload(text, name="data.csv"):
    buffer = io.BytesIO(text.encode("utf-8"))
    buffer.name = name
    return buffer


class FakeUpload(io.BytesIO):
    def __init__(self, data, name):
        super().__init__(data)
        self.name, self.size = name, len(data)


# --- reading and writing tables ---------------------------------------------

def test_read_table_handles_csv_and_excel():
    csv_rows = read_table(FakeUpload("﻿Name,Score\nAma,12\n,\n".encode(), "x.csv"), max_bytes=10_000)
    assert csv_rows == [{"name": "Ama", "score": "12"}]
    data = xlsx_upload([["Name", "Score", "Born"], ["Kofi", 12.0, None]]).getvalue()
    assert read_table(FakeUpload(data, "x.xlsx"), max_bytes=100_000) == [{"name": "Kofi", "score": "12", "born": ""}]
    with pytest.raises(ValueError, match="xlsx"):
        read_table(FakeUpload(b"junk", "old.xls"), max_bytes=10_000)
    with pytest.raises(ValueError, match="too large"):
        read_table(FakeUpload(b"a" * 20, "x.csv"), max_bytes=10)


def test_exports_in_excel(world, registrar):
    world.student.first_name = "=1+1"
    world.student.save()
    response = client_for(registrar).get("/api/v1/students/export/?file_format=xlsx")
    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    sheet = load_workbook(io.BytesIO(response.content)).active
    values = [cell.value for cell in sheet[2]]
    assert sheet["A1"].value == "student_number"
    assert "'=1+1" in values  # formulas are never written as formulas
    template = client_for(registrar).get("/api/v1/students/import/template/?file_format=xlsx")
    assert template["Content-Disposition"].endswith('.xlsx"')


def test_student_import_accepts_excel(world, registrar):
    upload = xlsx_upload([
        ["first_name", "last_name", "email", "programme_code", "level", "admission_year", "date_of_birth"],
        ["Yaa", "Badu", "yaa.badu@example.test", "BSC-CS", "Level 100", 2025, None],
    ])
    response = client_for(registrar).post("/api/v1/students/import/preview/", {"file": upload}, format="multipart")
    assert response.status_code == 200, response.data
    assert response.data["valid"] == 1


# --- marks import -----------------------------------------------------------

@pytest.fixture
def marked_sheet(world):
    for student in (world.student, world.other_student):
        registration.submit_registration(student, world.semester,
                                         [world.offerings["CS101"].pk, world.offerings["MATH101"].pk], actor=student.user)
    sheet = results.get_or_create_sheet(world.offerings["CS101"])
    results.replace_components(sheet, [
        {"name": "Coursework", "weight": 40, "max_score": 40},
        {"name": "Examination", "weight": 60, "max_score": 100, "is_examination": True},
    ], actor=world.lecturer)
    return sheet


def test_marks_template_lists_class_and_components(world, marked_sheet):
    response = client_for(world.lecturer).get(f"/api/v1/result-sheets/{marked_sheet.id}/marks-template/?file_format=xlsx")
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert [c.value for c in sheet[1]] == ["student_number", "student_name", "Coursework (/40)", "Examination (/100)"]
    assert sheet.max_row == 3


def test_marks_import_preview_then_apply(world, marked_sheet):
    api = client_for(world.lecturer)
    upload = xlsx_upload([
        ["Student_Number", "Student Name", "Coursework (/40)", "EXAMINATION", "Notes"],
        [world.student.student_number.lower(), "x", 35, 80, "ignored"],
        [world.other_student.student_number, "y", 45, "", ""],  # 45 > 40
        ["NOT-A-STUDENT", "z", 10, 10, ""],
    ])
    preview = api.post(f"/api/v1/result-sheets/{marked_sheet.id}/marks-import/", {"file": upload}, format="multipart")
    assert preview.status_code == 200, preview.data
    data = preview.data
    assert data["valid_rows"] == 1 and data["invalid_rows"] == 2
    assert data["ignored_columns"] == ["notes"]
    assert len(data["entries"]) == 2
    assert not StudentMark.objects.exists()  # preview never saves

    applied = api.put(f"/api/v1/result-sheets/{marked_sheet.id}/marks/",
                      {"marks": data["entries"], "source": "import"}, format="json")
    assert applied.status_code == 200 and applied.data["changed"] == 2
    log = AuditLog.objects.get(action="marks.saved")
    assert log.metadata["source"] == "import"

    # Re-importing the same values changes nothing.
    again = api.post(f"/api/v1/result-sheets/{marked_sheet.id}/marks-import/",
                     {"file": csv_upload(f"student_number,Coursework,Examination\n{world.student.student_number},35,80\n")},
                     format="multipart")
    assert again.data["entries"] == []


def test_marks_import_permissions_and_locking(world, marked_sheet):
    upload = csv_upload(f"student_number,Coursework\n{world.student.student_number},10\n")
    stranger = make_user(Role.LECTURER)
    assert client_for(stranger).post(f"/api/v1/result-sheets/{marked_sheet.id}/marks-import/", {"file": upload},
                                     format="multipart").status_code == 404
    officer = make_user(Role.RESULTS_OFFICER)
    upload.seek(0)
    assert client_for(officer).post(f"/api/v1/result-sheets/{marked_sheet.id}/marks-import/", {"file": upload},
                                    format="multipart").status_code == 403
    marked_sheet.status = "SUBMITTED"
    marked_sheet.save()
    upload.seek(0)
    assert client_for(world.lecturer).post(f"/api/v1/result-sheets/{marked_sheet.id}/marks-import/", {"file": upload},
                                           format="multipart").status_code == 409


# --- course import ------------------------------------------------------------

def test_course_import_with_prerequisites_and_curriculum(world, admin_user):
    api = client_for(admin_user)
    text = ("code,title,credits,department_code,course_type,prerequisites,programme_code,level,term_number,category\n"
            "cs301,Algorithms,3,cs,theory,CS201;CS101,BSC-CS,Level 200,1,elective\n"
            "CS302,Compilers,3,CS,,CS301,,,,\n"
            "CS303,Loop A,3,CS,,CS304,,,,\n"
            "CS304,Loop B,3,CS,,CS303,,,,\n"
            "CS101,Duplicate,3,CS,,,,,,\n"
            "CS305,Bad level,3,CS,,,BSC-CS,,,\n")
    preview = api.post("/api/v1/courses/import/preview/", {"file": csv_upload(text)}, format="multipart")
    assert preview.status_code == 200, preview.data
    rows = {r["input"]["code"]: r for r in preview.data["rows"]}
    assert rows["cs301"]["valid"] and rows["CS302"]["valid"]
    assert "circular" in str(rows["CS303"]["errors"]) and "already exists" in str(rows["CS101"]["errors"])
    assert "level" in rows["CS305"]["errors"]

    rejected = api.post("/api/v1/courses/import/confirm/", {"rows": [r["input"] for r in preview.data["rows"]]}, format="json")
    assert rejected.status_code == 400 and not Course.objects.filter(code="CS301").exists()

    good = [rows["cs301"]["input"], rows["CS302"]["input"]]
    created = api.post("/api/v1/courses/import/confirm/", {"rows": good}, format="json")
    assert created.status_code == 201 and created.data == {"created": 2, "placements": 1}
    algorithms = Course.objects.get(code="CS301")
    assert set(algorithms.prerequisites.values_list("code", flat=True)) == {"CS101", "CS201"}
    assert ProgrammeCourse.objects.get(course=algorithms).category == "ELECTIVE"
    assert list(Course.objects.get(code="CS302").prerequisites.values_list("code", flat=True)) == ["CS301"]
    assert client_for(make_user(Role.REGISTRAR)).post("/api/v1/courses/import/confirm/", {"rows": good},
                                                      format="json").status_code == 403


# --- staff import -------------------------------------------------------------

def test_staff_import_sends_setup_links_without_passwords(admin_user, api):
    client = client_for(admin_user)
    upload = xlsx_upload([
        ["email", "title", "first_name", "last_name", "staff_id", "role"],
        ["New.Lecturer@Example.test", "Dr.", "Ama", "Nkansah", "STF-9", "Lecturer"],
        ["officer@example.test", "", "Kojo", "Mensah", "", "Exams officer"],
        ["bad", "", "X", "Y", "", "STUDENT"],
    ])
    preview = client.post("/api/v1/users/import/preview/", {"file": upload}, format="multipart")
    assert preview.status_code == 200 and preview.data["valid"] == 2
    bad = preview.data["rows"][2]["errors"]
    assert "email" in bad and "role" in bad  # students cannot be created here

    rows = [r["input"] for r in preview.data["rows"] if r["valid"]]
    with __import__("django.test", fromlist=["TestCase"]).TestCase.captureOnCommitCallbacks(execute=True):
        created = client.post("/api/v1/users/import/confirm/", {"rows": rows}, format="json")
    assert created.status_code == 201 and created.data["created"] == 2
    lecturer = User.objects.get(email="new.lecturer@example.test")
    officer = User.objects.get(email="officer@example.test")
    assert lecturer.role == Role.LECTURER and not lecturer.has_usable_password()
    assert officer.role == Role.RESULTS_OFFICER and officer.can_approve_results
    assert len(mail.outbox) == 2
    uid, token = re.search(r"uid=([^&\s]+)&token=([^\s]+)", mail.outbox[0].body).groups()
    done = api.post("/api/v1/auth/password/reset/confirm/", {"uid": uid, "token": token, "new_password": "Brand-New-Pass-1"},
                    format="json")
    assert done.status_code == 204


def test_forgot_password_works_for_accounts_without_password(world, api):
    world.student.user.set_unusable_password()
    world.student.user.save()
    api.post("/api/v1/auth/password/reset/", {"email": world.student.email}, format="json")
    assert len(mail.outbox) == 1


# --- PDF documents --------------------------------------------------------------

def test_pdf_documents(world, registrar):
    reg = registration.submit_registration(world.student, world.semester,
                                           [world.offerings["CS101"].pk, world.offerings["MATH101"].pk],
                                           actor=world.student.user)
    other = registration.submit_registration(world.other_student, world.semester,
                                             [world.offerings["CS101"].pk, world.offerings["MATH101"].pk],
                                             actor=world.other_student.user)
    student_api = client_for(world.student.user)
    slip = student_api.get(f"/api/v1/registrations/{reg.id}/slip-pdf/")
    assert slip.status_code == 200 and slip["Content-Type"] == "application/pdf" and slip.content[:4] == b"%PDF"
    assert student_api.get(f"/api/v1/registrations/{other.id}/slip-pdf/").status_code == 404
    assert student_api.get("/api/v1/my-results/transcript-pdf/").content[:4] == b"%PDF"
    assert student_api.get(f"/api/v1/students/{world.student.pk}/transcript-pdf/").status_code == 403

    staff = client_for(registrar).get(f"/api/v1/students/{world.student.pk}/transcript-pdf/?published_only=false")
    assert staff.status_code == 200 and staff.content[:4] == b"%PDF"
    assert AuditLog.objects.filter(action="student.transcript_generated").exists()
