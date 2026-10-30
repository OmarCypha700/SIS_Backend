"""
Spreadsheet input and output shared by every import and export: CSV (UTF-8) and Excel (.xlsx).

Exports choose the format with the `file_format` query parameter (`csv`, the default, or `xlsx`).
`format` is not used because Django REST framework reserves it for content negotiation.
"""
import csv
import datetime
import io
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMATS = ("csv", "xlsx")

# Characters that spreadsheet applications interpret as the start of a formula.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    """Text for a cell; values that look like formulas are escaped so they are never executed."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, (int, float, Decimal)):
        return value
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.isoformat()
    text = str(value)
    if text.startswith(_FORMULA_PREFIXES) and not _is_number(text):
        return "'" + text
    return text


def _is_number(text):
    try:
        float(text)
        return True
    except ValueError:
        return False


def requested_format(request):
    value = (request.query_params.get("file_format") or "csv").lower()
    return value if value in FORMATS else "csv"


def table_response(filename_base, header, rows, file_format="csv", sheet_title="Data"):
    """Download response for a header + rows table in CSV or Excel."""
    if file_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_title[:31]
        sheet.append([safe_cell(str(h)) for h in header])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        widths = [len(str(h)) for h in header]
        for row in rows:
            values = [safe_cell(v) for v in row]
            sheet.append(values)
            for i, value in enumerate(values[: len(widths)]):
                widths[i] = max(widths[i], min(len(str(value)), 60))
        for i, width in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=i).column_letter].width = width + 2
        sheet.freeze_panes = "A2"
        buffer = io.BytesIO()
        workbook.save(buffer)
        response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
        response["Content-Disposition"] = f'attachment; filename="{filename_base}.xlsx"'
        return response

    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename_base}.csv"'
    response.write("﻿")  # BOM so Excel opens UTF-8 CSV files correctly
    writer = csv.writer(response)
    writer.writerow([safe_cell(str(h)) for h in header])
    for row in rows:
        writer.writerow([safe_cell(v) for v in row])
    return response


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.date().isoformat() if value.time() == datetime.time() else value.isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


MAX_COLUMNS = 100


def read_table(uploaded_file, *, max_bytes, max_rows=5000):
    """
    Read an uploaded CSV or .xlsx file into a list of dicts keyed by lower-case header names.
    Rows are read one at a time and reading stops as soon as a limit is passed, so a small file
    that declares a huge sheet cannot exhaust memory. Raises ValueError with a user-facing message.
    """
    if uploaded_file.size > max_bytes:
        raise ValueError(f"File is too large (maximum {max_bytes // (1024 * 1024)} MB).")
    name = (uploaded_file.name or "").lower()
    raw = uploaded_file.read()
    if name.endswith(".xls") and raw[:2] != b"PK":
        raise ValueError("Old .xls files are not supported. Save the file as .xlsx or .csv.")
    too_many = ValueError(f"Import at most {max_rows} rows at a time.")
    # Blank lines are skipped, but every line read counts towards this hard stop.
    line_limit = max_rows * 2 + 50
    rows = _iter_xlsx(raw) if name.endswith(".xlsx") or raw[:2] == b"PK" else _iter_csv(raw)
    header, records = None, []
    try:
        for count, values in enumerate(rows, start=1):
            if count > line_limit:
                raise too_many
            values = [_cell_text(v) for v in values[:MAX_COLUMNS]]
            if not any(values):
                continue  # skip blank lines (and empty rows above the header)
            if header is None:
                header = [h.lower() for h in values]
                continue
            records.append({key: (values[i] if i < len(values) else "") for i, key in enumerate(header) if key})
            if len(records) > max_rows:
                raise too_many
    finally:
        close = getattr(rows, "close", None)
        if close:
            close()
    if not header:
        raise ValueError("The file has no header row.")
    return records


def _iter_csv(raw):
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV files must be saved as UTF-8.") from exc
    yield from csv.reader(io.StringIO(text))


def _iter_xlsx(raw):
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise ValueError("This is not a valid Excel (.xlsx) file.") from exc
    try:
        sheet = workbook.worksheets[0]
        # Ignore the sheet's declared size, which the file can set to anything.
        sheet.reset_dimensions()
        for row in sheet.iter_rows(values_only=True, max_col=MAX_COLUMNS):
            yield list(row)
    finally:
        workbook.close()
